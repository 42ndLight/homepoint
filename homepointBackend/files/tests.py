import io
import json
from unittest.mock import patch
import openpyxl
from django.test import TestCase
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from products.models import Product, Variant, Category, Inventory
from files.models import ImportHistory, PendingUpload
from files.tasks import process_xlsx_import_task, reconcile_stalled_uploads
from django.utils import timezone
from datetime import timedelta

class FilesAppTests(TestCase):
    def setUp(self):
        # Create a category to reference
        self.category = Category.objects.create(name="Pipes", slug="pipes")
        
        # Create a dummy xlsx file
        wb = openpyxl.Workbook()
        
        # Categories sheet
        ws_cat = wb.active
        ws_cat.title = "Categories"
        ws_cat.append(["id", "name", "slug", "description", "parent id"])
        ws_cat.append(["Notes row"])
        ws_cat.append([1, "Pipes", "pipes", "Pipe description", None])
        
        # Products sheet
        ws_prod = wb.create_sheet("Products")
        ws_prod.append(["product id", "name", "slug", "description", "category id", "base price (kes)", "is active"])
        ws_prod.append(["Notes row"])
        ws_prod.append([1, "PVC Pipe", "pvc-pipe", "Desc", 1, 100, 1])

        # Variants sheet
        ws_var = wb.create_sheet("Variants")
        # Ensure we have "quantity" in the header to test the fallback!
        ws_var.append(["variant id", "product id", "sku", "price", "quantity", "unit type", "tax type", "item code", "low stock alert"])
        ws_var.append(["Notes row"])
        ws_var.append([101, 1, "SKU-101", 120, 50, "piece", "A", "CODE-1", 10])

        # Inventory sheet (empty or missing qty to test Variants fallback)
        ws_inv = wb.create_sheet("Inventory")
        ws_inv.append(["variant id", "quantity", "location", "last updated"])
        ws_inv.append(["Notes row"])
        ws_inv.append([101, None, "Warehouse A", ""]) # Quantity is None here!

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        self.file_path = "test_import.xlsx"
        default_storage.save(self.file_path, ContentFile(buf.read()))
        
        self.history = ImportHistory.objects.create(
            task_id="test-task-1",
            status="PENDING",
            file_path=self.file_path
        )

    def tearDown(self):
        if default_storage.exists(self.file_path):
            default_storage.delete(self.file_path)

    def test_xlsx_import_quantity(self):
        """Test that quantity is correctly read from the Variants sheet when importing."""
        # Use apply to run synchronously with a specific task_id
        process_xlsx_import_task.apply(args=[self.file_path], task_id="test-task-1")
        
        self.history.refresh_from_db()
        self.assertEqual(self.history.status, "COMPLETED")
        
        # Verify Product was created
        prod = Product.objects.get(slug="pvc-pipe")
        self.assertEqual(prod.name, "PVC Pipe")
        
        # Verify Variant was created
        var = Variant.objects.get(sku="SKU-101")
        self.assertEqual(var.price, 120)
        
        # Verify Inventory has the correct quantity (50 from Variants sheet)
        inv = var.inventory
        self.assertEqual(inv.quantity, 50)
        self.assertEqual(inv._qty if hasattr(inv, '_qty') else inv.quantity, 50)

    def test_reconcile_stalled_uploads(self):
        # Create a product to satisfy foreign key constraint
        Product.objects.create(id=1, name="Test Prod", category=self.category, base_price=10)
        
        # Create an expired pending upload where the file doesn't exist
        p1 = PendingUpload.objects.create(
            record_id="00000000-0000-0000-0000-000000000001",
            storage_key="raw/product_image/test_missing.jpg",
            kind="product_image",
            target_id="1",
            status="pending"
        )
        PendingUpload.objects.filter(id=p1.id).update(created_at=timezone.now() - timedelta(minutes=15))
        
        # Create a valid one where the file does exist
        valid_key = "raw/product_image/test_present.jpg"
        default_storage.save(valid_key, ContentFile(b"fake image data"))
        
        p2 = PendingUpload.objects.create(
            record_id="00000000-0000-0000-0000-000000000002",
            storage_key=valid_key,
            kind="product_image",
            target_id="1",
            status="pending"
        )
        PendingUpload.objects.filter(id=p2.id).update(created_at=timezone.now() - timedelta(minutes=15))
        
        # Also create a fresh one which shouldn't be touched
        PendingUpload.objects.create(
            record_id="00000000-0000-0000-0000-000000000003",
            storage_key="raw/product_image/fresh.jpg",
            kind="product_image",
            target_id="1",
            status="pending"
        )
        
        # We need to mock process_image_optimization_task.delay so we don't actually trigger it
        with patch('products.tasks.process_image_optimization_task.delay') as mock_delay:
            result = reconcile_stalled_uploads()
            
            self.assertEqual(result["expired"], 1)
            self.assertEqual(result["promoted"], 1)
            
            p1.refresh_from_db()
            self.assertEqual(p1.status, "expired")
            
            p2.refresh_from_db()
            self.assertEqual(p2.status, "uploaded")
            
            p3 = PendingUpload.objects.get(record_id="00000000-0000-0000-0000-000000000003")
            self.assertEqual(p3.status, "pending")
            
            mock_delay.assert_called_once()

        default_storage.delete(valid_key)
